# -*- coding: utf-8 -*-
"""
서버 연락처 / 수신 규칙 등록 시스템
- Python 3.13 / Flask / waitress / openpyxl
- WAS 와 같은 머신, Node.js(relayServer) 가 읽는 JSON 을 로컬 경로에 원자적 기록 (Node 는 읽기만)
- 로그인 + 권한(admin/viewer) / 서버등록 / 일괄등록 / 규칙 일괄적용 / 계정관리 / 변경이력
- 레코드: hostname, ip, department, rules[]
    rule: id, name, metrics[], keywords[], any_level(Y/N), DAY(Y/N), NIGHT(Y/N), contacts[{name, phone}]
    · metrics  : 관제 metric 이름 또는 '*' 와일드카드 (['*'] = 기본 규칙: 다른 규칙이 안 맞을 때만)
    · keywords : memo 포함 문자열 (있으면 등급 무관 발신)
    · contacts : 순서 = 정담당 → 부담당 … (각자 최대 3회, 미응답 시 다음 사람)
"""

import csv
import io
import json
import os
import re
import secrets
import tempfile
import threading
import math
from datetime import datetime
from functools import wraps

from flask import (
    Flask, render_template, request, redirect, url_for,
    session, flash, send_file, abort, Response, jsonify,
)
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook, load_workbook

# ===========================================================================
#  ▼▼▼ 설정 — 운영 환경에 맞게 이 블록만 수정하세요 ▼▼▼
# ===========================================================================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Node.js(relayServer) 와 공유하는 호스트 데이터 JSON 경로
#   기본값: <repo>/hostRegistry/data/hosts.json (relayServer adapter 기본값과 동일)
#   위치를 바꿀 때는 양쪽 모두 환경변수 HOST_DATA_FILE 로 같은 경로를 지정
HOST_DATA_FILE = os.environ.get(
    "HOST_DATA_FILE", os.path.join(BASE_DIR, "data", "hosts.json")
)

SERVER_HOST = os.environ.get("SERVER_HOST", "0.0.0.0")    # 바인딩 IP
SERVER_PORT = int(os.environ.get("SERVER_PORT", "8458"))  # Node 와 다른 포트로!

USERS_FILE = os.environ.get("USERS_FILE", os.path.join(BASE_DIR, "data", "users.json"))
AUDIT_FILE = os.environ.get("AUDIT_FILE", os.path.join(BASE_DIR, "data", "audit.jsonl"))
# 규칙 입력 시 metric 자동완성 목록 (한 줄에 하나)
METRIC_CATALOG_FILE = os.environ.get("METRIC_CATALOG_FILE",
                                     os.path.join(BASE_DIR, "data", "metric_catalog.txt"))

PAGE_SIZE = 50
ALLOWED_EXT = {".csv", ".xlsx"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
# ===========================================================================
#  ▲▲▲ 설정 끝 ▲▲▲
# ===========================================================================

SECRET_FILE = os.path.join(BASE_DIR, "data", "secret.key")

HOST_FIELDS = ["hostname", "ip", "department"]
CONTACT_FIELDS = ["name", "phone"]
FLAG_FIELDS = ["DAY", "NIGHT"]
CORE_FIELDS = HOST_FIELDS                      # (템플릿 호환)
SEARCH_FIELDS = HOST_FIELDS + CONTACT_FIELDS + ["rule", "metric"]
SORT_FIELDS = HOST_FIELDS + ["rules"]
FIELD_LABELS = {
    "hostname": "호스트네임", "ip": "IP", "department": "부서", "name": "담당자", "phone": "핸드폰번호",
    "DAY": "주간", "NIGHT": "야간", "rules": "수신 규칙", "rule": "규칙명", "metric": "metric",
    "keyword": "memo 키워드", "any_level": "등급 무관",
}
ACTION_LABELS = {
    "login": "로그인", "logout": "로그아웃", "add": "서버추가", "edit": "서버수정",
    "delete": "서버삭제", "import": "일괄등록", "rules_bulk": "규칙 일괄적용",
    "export_csv": "CSV 내보내기", "export_xlsx": "Excel 내보내기",
    "user_add": "계정생성", "user_role": "권한변경", "user_reset": "비번초기화",
    "user_delete": "계정삭제", "change_password": "비번변경",
}

# ===========================================================================
app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES
_hosts_lock = threading.Lock()
_users_lock = threading.Lock()
_audit_lock = threading.Lock()


def _load_or_create_secret():
    os.makedirs(os.path.dirname(SECRET_FILE), exist_ok=True)
    if os.path.exists(SECRET_FILE):
        with open(SECRET_FILE, "r", encoding="utf-8") as f:
            return f.read().strip()
    key = secrets.token_hex(32)
    with open(SECRET_FILE, "w", encoding="utf-8") as f:
        f.write(key)
    return key


app.secret_key = _load_or_create_secret()




# ===========================================================================
#  저장소
# ===========================================================================
def _atomic_write_json(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=os.path.dirname(path), suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
            f.flush(); os.fsync(f.fileno())
        os.replace(tmp, path)
    except Exception:
        if os.path.exists(tmp):
            os.unlink(tmp)
        raise


def _yn(v):
    return "Y" if str(v or "").strip().upper() == "Y" else "N"


def _to_list(v, sep=";"):
    if isinstance(v, list):
        return [str(x).strip() for x in v if str(x).strip()]
    if isinstance(v, str):
        return [x.strip() for x in re.split(rf"[{sep}\n]", v) if x.strip()]
    return []


def new_rule_id():
    return "r" + secrets.token_hex(3)


def _normalize_rule(r, idx=0):
    contacts = []
    for c in r.get("contacts", []) or []:
        phone = normalize_phone(str(c.get("phone", "")))
        if phone:
            contacts.append({"name": str(c.get("name", "") or "").strip(), "phone": phone})
    keywords = _to_list(r.get("keywords", r.get("keyword", [])))
    metrics = _to_list(r.get("metrics", [])) or ["*"]
    return {
        "id": str(r.get("id") or f"r{idx + 1}"),
        "name": str(r.get("name", "") or "").strip() or f"규칙{idx + 1}",
        "metrics": metrics,
        "keywords": keywords,
        "any_level": "Y" if keywords else _yn(r.get("any_level")),
        "DAY": _yn(r.get("DAY")),
        "NIGHT": _yn(r.get("NIGHT")),
        "contacts": contacts,
    }


def _normalize_record(r):
    """기존 단일 담당자 레코드(name/phone/DAY/NIGHT/keyword)는 '기본' 규칙 1개로 변환"""
    out = {f: str(r.get(f, "") or "").strip() for f in HOST_FIELDS}
    if isinstance(r.get("rules"), list):
        rules = r["rules"]
    else:
        kw = _to_list(r.get("keyword", []))
        rules = [{"id": "default", "name": "기본", "metrics": ["*"], "keywords": kw,
                  "any_level": "Y" if kw else "N", "DAY": r.get("DAY"), "NIGHT": r.get("NIGHT"),
                  "contacts": [{"name": r.get("name", ""), "phone": r.get("phone", "")}]}]
    out["rules"] = [_normalize_rule(x, i) for i, x in enumerate(rules)]
    return out


def load_hosts():
    if not os.path.exists(HOST_DATA_FILE):
        return []
    try:
        with open(HOST_DATA_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        return [_normalize_record(r) for r in data] if isinstance(data, list) else []
    except (json.JSONDecodeError, OSError):
        return []


def save_hosts(data):
    _atomic_write_json(HOST_DATA_FILE, data)


def migrate_existing():
    """기존 hosts.json(서버당 담당자 1명) → 규칙 구조로 1회 변환. 원본은 .bak 으로 보관"""
    if not os.path.exists(HOST_DATA_FILE):
        return
    try:
        with open(HOST_DATA_FILE, "r", encoding="utf-8") as f:
            raw = json.load(f)
    except (json.JSONDecodeError, OSError):
        return
    if isinstance(raw, list) and any(not isinstance(r.get("rules"), list) for r in raw):
        with _hosts_lock:
            backup = f"{HOST_DATA_FILE}.{datetime.now():%Y%m%d_%H%M%S}.bak"
            _atomic_write_json(backup, raw)
            save_hosts([_normalize_record(r) for r in raw])


def load_metric_catalog():
    try:
        with open(METRIC_CATALOG_FILE, "r", encoding="utf-8") as f:
            return [line.strip() for line in f if line.strip()]
    except OSError:
        return []


def load_users():
    if not os.path.exists(USERS_FILE):
        return []
    try:
        with open(USERS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except (json.JSONDecodeError, OSError):
        return []


def save_users(users):
    _atomic_write_json(USERS_FILE, users)


def _now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def ensure_default_admin():
    with _users_lock:
        users = load_users()
        if not any(u.get("username") == "admin" for u in users):
            users.append({"username": "admin", "password_hash": generate_password_hash("admin"),
                          "role": "admin", "is_default": True, "created_at": _now()})
            save_users(users)


def audit(action, target="", detail=None):
    try:
        ip = request.remote_addr or "-"
    except Exception:
        ip = "-"
    entry = {"ts": _now(), "ip": ip, "user": session.get("username", "-"),
             "action": action, "target": target, "detail": detail}
    with _audit_lock:
        os.makedirs(os.path.dirname(AUDIT_FILE), exist_ok=True)
        with open(AUDIT_FILE, "a", encoding="utf-8") as f:
            f.write(json.dumps(entry, ensure_ascii=False) + "\n")



# ===========================================================================
#  검증 / 조회
# ===========================================================================
PHONE_RE = re.compile(r"^010\d{8}$")
IP_RE = re.compile(r"^(\d{1,3})\.(\d{1,3})\.(\d{1,3})\.(\d{1,3})$")


def normalize_phone(raw):
    return re.sub(r"\D", "", raw or "")


def is_valid_phone(d):
    return bool(PHONE_RE.match(d))


def is_valid_ip(ip):
    m = IP_RE.match(ip or "")
    return bool(m) and all(0 <= int(x) <= 255 for x in m.groups())


def find_host(hosts, hostname):
    key = (hostname or "").strip().lower()
    return next((h for h in hosts if h.get("hostname", "").strip().lower() == key), None)


def find_by_ip(hosts, ip, exclude_hostname=None):
    ip = (ip or "").strip()
    excl = (exclude_hostname or "").strip().lower()
    if not ip:
        return None
    for h in hosts:
        if h.get("hostname", "").strip().lower() == excl:
            continue
        if h.get("ip", "").strip() == ip:
            return h
    return None


def all_contacts(host):
    for r in host.get("rules", []):
        for c in r.get("contacts", []):
            yield r, c


def contact_conflicts(hosts, rules, exclude_hostname=None):
    """같은 이름인데 다른 번호로 등록된 담당자 (오타 확인용 경고)"""
    excl = (exclude_hostname or "").strip().lower()
    known = {}
    for h in hosts:
        if h.get("hostname", "").strip().lower() == excl:
            continue
        for _, c in all_contacts(h):
            if c["name"]:
                known.setdefault(c["name"], {}).setdefault(c["phone"], h["hostname"])
    out = []
    for r in rules:
        for c in r["contacts"]:
            others = {p: hn for p, hn in known.get(c["name"], {}).items() if p != c["phone"]}
            for p, hn in others.items():
                out.append({"name": c["name"], "phone": c["phone"], "other_phone": p, "other_host": hn})
    return out


def validate_host(rec):
    errors, out = [], {}
    for f in HOST_FIELDS:
        out[f] = (rec.get(f) or "").strip()
        if not out[f]:
            errors.append(f"{FIELD_LABELS[f]} 누락")
    if out["ip"] and not is_valid_ip(out["ip"]):
        errors.append("IP 형식 오류(IPv4)")
    return errors, out


def validate_rules(raw_rules, prefix=""):
    """규칙 목록 검증 → (errors, 정규화된 rules)"""
    errors, rules, names = [], [], set()
    if not isinstance(raw_rules, list) or not raw_rules:
        return [f"{prefix}수신 규칙이 최소 1개 필요합니다."], []
    for i, r in enumerate(raw_rules):
        label = f"{prefix}규칙 {i + 1}"
        name = str(r.get("name", "") or "").strip()
        if not name:
            errors.append(f"{label}: 규칙명 누락")
        elif name.lower() in names:
            errors.append(f"{label}: 규칙명 '{name}' 중복")
        names.add(name.lower())
        contacts = []
        for j, c in enumerate(r.get("contacts", []) or []):
            cname = str(c.get("name", "") or "").strip()
            phone = normalize_phone(str(c.get("phone", "") or ""))
            if not cname and not phone:
                continue
            if not cname:
                errors.append(f"{label} 담당자 {j + 1}: 이름 누락")
            if not is_valid_phone(phone):
                errors.append(f"{label} 담당자 {j + 1}: 핸드폰번호 형식 오류(010으로 시작하는 11자리)")
            contacts.append({"name": cname, "phone": phone})
        if not contacts:
            errors.append(f"{label}: 담당자가 최소 1명 필요합니다.")
        nr = _normalize_rule({**r, "contacts": contacts}, i)
        nr["contacts"] = contacts
        if not r.get("id"):
            nr["id"] = new_rule_id()
        rules.append(nr)
    return errors, rules


def is_catch_all(rule):
    return all(m.strip() in ("", "*") for m in rule["metrics"]) and not rule["keywords"]


def host_flags(host):
    rules = host.get("rules", [])
    return {"DAY": "Y" if any(r["DAY"] == "Y" for r in rules) else "N",
            "NIGHT": "Y" if any(r["NIGHT"] == "Y" for r in rules) else "N"}


def sort_hosts(hosts, sort, direction):
    if sort not in SORT_FIELDS:
        sort = "hostname"
    rev = (direction == "desc")
    if sort == "rules":
        key = lambda h: len(h.get("rules", []))
    elif sort == "ip":
        def key(h):
            try:
                return tuple(int(x) for x in h.get("ip", "").split("."))
            except (ValueError, AttributeError):
                return (0,)
    else:
        key = lambda h: str(h.get(sort, "")).lower()
    return sorted(hosts, key=key, reverse=rev)


def summary(hosts):
    return {"total": len(hosts),
            "multi": sum(1 for h in hosts if len(h.get("rules", [])) > 1),
            "contacts": len({c["phone"] for h in hosts for _, c in all_contacts(h)}),
            "day": sum(1 for h in hosts if host_flags(h)["DAY"] == "Y"),
            "night": sum(1 for h in hosts if host_flags(h)["NIGHT"] == "Y")}


def rule_text(r):
    flags = "/".join(x for x, f in (("주간", r["DAY"]), ("야간", r["NIGHT"])) if f == "Y") or "발신 안 함"
    cond = ", ".join(r["metrics"])
    if r["keywords"]:
        cond += f" + memo[{', '.join(r['keywords'])}]"
    who = " → ".join(f"{c['name']}({c['phone']})" for c in r["contacts"])
    level = "등급무관" if r["any_level"] == "Y" else "심각만"
    return f"[{r['name']}] {cond} · {flags} · {level} · {who}"


def record_detail(rec):
    d = {FIELD_LABELS[f]: rec.get(f, "") for f in HOST_FIELDS}
    d[FIELD_LABELS["rules"]] = [rule_text(r) for r in rec.get("rules", [])]
    return d


def distinct_values(hosts, field):
    return sorted({h.get(field, "").strip() for h in hosts if h.get(field, "").strip()})


def distinct_contacts(hosts):
    seen = {}
    for h in hosts:
        for _, c in all_contacts(h):
            if c["name"]:
                seen.setdefault((c["name"], c["phone"]), None)
    return [{"name": n, "phone": p} for n, p in sorted(seen)]


def paginate(items, page):
    pages = max(1, math.ceil(len(items) / PAGE_SIZE))
    page = max(1, min(page, pages))
    s = (page - 1) * PAGE_SIZE
    return items[s:s + PAGE_SIZE], page, pages, len(items)


def _page_arg():
    try:
        return int(request.args.get("page", 1) or 1)
    except ValueError:
        return 1


def host_matches(h, field, kw):
    kw = kw.lower()
    if field in HOST_FIELDS:
        return kw in str(h.get(field, "")).lower()
    if field in CONTACT_FIELDS:
        if field == "phone":
            kw = normalize_phone(kw) or kw
        return any(kw in str(c.get(field, "")).lower() for _, c in all_contacts(h))
    if field == "rule":
        return any(kw in r["name"].lower() for r in h.get("rules", []))
    if field == "metric":
        return any(kw in m.lower() for r in h.get("rules", []) for m in r["metrics"] + r["keywords"])
    return any(host_matches(h, f, kw) for f in SEARCH_FIELDS)


# ===========================================================================
#  인증
# ===========================================================================
def login_required(f):
    @wraps(f)
    def w(*a, **k):
        if not session.get("username"):
            return redirect(url_for("login", next=request.path))
        return f(*a, **k)
    return w


def admin_required(f):
    @wraps(f)
    def w(*a, **k):
        if not session.get("username"):
            return redirect(url_for("login", next=request.path))
        if session.get("role") != "admin":
            abort(403)
        return f(*a, **k)
    return w


@app.context_processor
def inject():
    return {"current_user": session.get("username"), "current_role": session.get("role"),
            "CORE_FIELDS": CORE_FIELDS, "FLAG_FIELDS": FLAG_FIELDS,
            "FIELD_LABELS": FIELD_LABELS, "SEARCH_FIELDS": SEARCH_FIELDS}


# ===========================================================================
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        u = next((x for x in load_users()
                  if x.get("username") == request.form.get("username", "").strip()), None)
        if u and check_password_hash(u["password_hash"], request.form.get("password", "")):
            session["username"] = u["username"]
            session["role"] = u["role"]
            session["is_default_pw"] = bool(u.get("is_default"))
            audit("login")
            return redirect(request.args.get("next") or url_for("index"))
        flash("아이디 또는 비밀번호가 올바르지 않습니다.", "error")
    return render_template("login.html")


@app.route("/logout")
def logout():
    audit("logout")
    session.clear()
    return redirect(url_for("login"))



# ===========================================================================
#  서버등록 목록
# ===========================================================================
def _search(hosts, field, kw):
    if not kw:
        return hosts
    return [h for h in hosts if host_matches(h, field, kw)]


def _rule_filter(hosts, pat):
    if pat == "multi":
        return [h for h in hosts if len(h.get("rules", [])) > 1]
    if pat == "single":
        return [h for h in hosts if len(h.get("rules", [])) <= 1]
    return hosts


def _search_args():
    field = request.args.get("field", "all").strip()
    if field not in SEARCH_FIELDS:
        field = "all"
    return field, request.args.get("q", "").strip(), request.args.get("pat", "all").strip()


@app.route("/")
@login_required
def index():
    field, q, pat = _search_args()
    sort = request.args.get("sort", "hostname")
    direction = request.args.get("dir", "asc")
    allh = load_hosts()
    hosts = sort_hosts(_rule_filter(_search(allh, field, q), pat), sort, direction)
    rows, page, pages, total = paginate(hosts, _page_arg())
    return render_template("index.html", hosts=rows, field=field, q=q, pat=pat, sort=sort, dir=direction,
                           page=page, pages=pages, total=total, stats=summary(allh))


# ===========================================================================
#  서버 추가 / 수정 / 삭제
# ===========================================================================
def _dup_errors(hosts, norm, exclude_hostname=None):
    errs = []
    if find_host(hosts, norm["hostname"]) and (exclude_hostname is None or
            norm["hostname"].strip().lower() != exclude_hostname.strip().lower()):
        errs.append(f"호스트네임 '{norm['hostname']}' 은(는) 이미 등록되어 있습니다.")
    ip_owner = find_by_ip(hosts, norm["ip"], exclude_hostname=exclude_hostname)
    if ip_owner:
        errs.append(f"IP '{norm['ip']}' 은(는) 이미 '{ip_owner['hostname']}' 에 등록되어 있습니다.")
    return errs


def _form(mode, rec, errors, warnings, original_hostname, hosts):
    return render_template("form.html", mode=mode, rec=rec, errors=errors, warnings=warnings,
                           original_hostname=original_hostname,
                           departments=distinct_values(hosts, "department"),
                           contacts=distinct_contacts(hosts), metrics=load_metric_catalog())


def _posted_host():
    rec = {f: request.form.get(f, "") for f in HOST_FIELDS}
    try:
        raw_rules = json.loads(request.form.get("rules_json", "[]") or "[]")
    except json.JSONDecodeError:
        raw_rules = None
    return rec, raw_rules


def _blank_rule():
    return {"id": "", "name": "기본", "metrics": ["*"], "keywords": [], "any_level": "N",
            "DAY": "Y", "NIGHT": "Y", "contacts": [{"name": "", "phone": ""}]}


def _save_host(mode, original_hostname=""):
    """add/edit POST 공통 처리. 성공 시 (None, out, before) / 화면 재표시 시 (response, None, None)"""
    rec, raw_rules = _posted_host()
    confirm = request.form.get("confirm_override") == "1"
    errors, norm = validate_host(rec)
    if raw_rules is None:
        errors.append("규칙 데이터를 해석할 수 없습니다.")
        raw_rules = []
    rule_errors, rules = validate_rules(raw_rules)
    errors += rule_errors
    norm["rules"] = rules or [_normalize_rule(r, i) for i, r in enumerate(raw_rules)]
    with _hosts_lock:
        hosts = load_hosts()
        before = None
        if mode == "edit":
            target = find_host(hosts, original_hostname)
            if not target:
                abort(404)
            before = json.loads(json.dumps(target))
        if not errors:
            errors += _dup_errors(hosts, norm, exclude_hostname=original_hostname or None)
        if errors:
            return _form(mode, norm, errors, [], original_hostname, hosts), None, None
        conflicts = contact_conflicts(hosts, rules, exclude_hostname=original_hostname or None)
        if conflicts and not confirm:
            return _form(mode, norm, [], conflicts, original_hostname, hosts), None, None
        out = {f: norm[f] for f in HOST_FIELDS}
        out["rules"] = rules
        if mode == "edit":
            hosts = [out if h is target else h for h in hosts]
        else:
            hosts.append(out)
        save_hosts(hosts)
    return None, out, before


@app.route("/add", methods=["GET", "POST"])
@admin_required
def add():
    if request.method == "POST":
        resp, out, _ = _save_host("add")
        if resp is not None:
            return resp
        audit("add", out["hostname"], record_detail(out))
        flash(f"'{out['hostname']}' 등록 완료.", "success")
        return redirect(url_for("index"))
    blank = {f: "" for f in HOST_FIELDS}
    blank["rules"] = [_blank_rule()]
    return _form("add", blank, [], [], "", load_hosts())


@app.route("/edit/<path:hostname>", methods=["GET", "POST"])
@admin_required
def edit(hostname):
    target = find_host(load_hosts(), hostname)
    if not target:
        abort(404)
    original_hostname = target["hostname"]
    if request.method == "POST":
        resp, out, before = _save_host("edit", original_hostname)
        if resp is not None:
            return resp
        b, a = record_detail(before), record_detail(out)
        changed = {k: [b[k], a[k]] for k in a if b.get(k) != a[k]}
        audit("edit", original_hostname, {"changed": changed})
        flash(f"'{out['hostname']}' 수정 완료.", "success")
        return redirect(url_for("index"))
    return _form("edit", target, [], [], original_hostname, load_hosts())


@app.route("/delete/<path:hostname>", methods=["POST"])
@admin_required
def delete(hostname):
    with _hosts_lock:
        hosts = load_hosts()
        target = find_host(hosts, hostname)
        if not target:
            abort(404)
        snapshot = dict(target)
        hosts = [h for h in hosts if h is not target]
        save_hosts(hosts)
    audit("delete", target["hostname"], record_detail(snapshot))
    flash(f"'{target['hostname']}' 삭제 완료.", "success")
    return redirect(url_for("index"))


# ===========================================================================
#  Export / Import 공통 (한 행 = 서버 × 규칙 × 담당자)
# ===========================================================================
IMPORT_HEADER = ["hostname", "ip", "department", "rule_name", "metrics", "keywords",
                 "any_level", "DAY", "NIGHT", "contact_seq", "name", "phone"]
IMPORT_REQUIRED = ["hostname", "ip", "department", "name", "phone"]
IMPORT_SAMPLE = [
    ["hl_rec_db", "10.5.240.5", "금융운영팀", "OS", "가용성;*사용률", "", "N", "Y", "Y", "1", "홍기웅", "010-1234-5678"],
    ["hl_rec_db", "10.5.240.5", "금융운영팀", "OS", "가용성;*사용률", "", "N", "Y", "Y", "2", "부담당", "010-2345-6789"],
    ["hl_rec_db", "10.5.240.5", "금융운영팀", "DBA", "SQL Server Error Log;*프로세스 개수", "", "Y", "Y", "Y", "1", "DBA담당", "010-3456-7890"],
    ["hl_rec_db", "10.5.240.5", "금융운영팀", "기본", "*", "", "N", "N", "Y", "1", "홍기웅", "010-1234-5678"],
]


def _export_rows(hosts):
    rows = [IMPORT_HEADER]
    for h in hosts:
        for r in h.get("rules", []):
            for seq, c in enumerate(r["contacts"], start=1):
                rows.append([h["hostname"], h["ip"], h["department"], r["name"], ";".join(r["metrics"]),
                             ";".join(r["keywords"]), r["any_level"], r["DAY"], r["NIGHT"], str(seq),
                             c["name"], c["phone"]])
    return rows


def _hosts_for_export():
    field, q, pat = _search_args()
    hosts = _rule_filter(_search(load_hosts(), field, q), pat)
    return sort_hosts(hosts, request.args.get("sort", "hostname"), request.args.get("dir", "asc"))


def _phone_text_cols(ws, header):
    col = header.index("phone") + 1
    for row in ws.iter_rows(min_col=col, max_col=col):
        for cell in row:
            cell.number_format = "@"


@app.route("/export/csv")
@login_required
def export_csv():
    hosts = _hosts_for_export()
    buf = io.StringIO(); w = csv.writer(buf)
    for r in _export_rows(hosts):
        w.writerow(r)
    audit("export_csv", f"{len(hosts)}대")
    fname = f"hosts_{datetime.now():%Y%m%d_%H%M%S}.csv"
    return Response(("﻿" + buf.getvalue()).encode("utf-8"), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={fname}"})


@app.route("/export/xlsx")
@login_required
def export_xlsx():
    hosts = _hosts_for_export()
    wb = Workbook(); ws = wb.active; ws.title = "hosts"
    for r in _export_rows(hosts):
        ws.append(r)
    _phone_text_cols(ws, IMPORT_HEADER)
    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    audit("export_xlsx", f"{len(hosts)}대")
    fname = f"hosts_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    return send_file(bio, as_attachment=True, download_name=fname,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/template/csv")
@login_required
def template_csv():
    buf = io.StringIO(); w = csv.writer(buf)
    w.writerow(IMPORT_HEADER)
    for r in IMPORT_SAMPLE:
        w.writerow(r)
    return Response(("﻿" + buf.getvalue()).encode("utf-8"), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=import_template.csv"})


@app.route("/template/xlsx")
@login_required
def template_xlsx():
    wb = Workbook(); ws = wb.active; ws.title = "hosts"
    ws.append(IMPORT_HEADER)
    for r in IMPORT_SAMPLE:
        ws.append(r)
    _phone_text_cols(ws, IMPORT_HEADER)
    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio, as_attachment=True, download_name="import_template.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _read_csv_rows(raw):
    for enc in ("utf-8-sig", "cp949"):
        try:
            return [dict(r) for r in csv.DictReader(io.StringIO(raw.decode(enc)))], None
        except (UnicodeDecodeError, csv.Error):
            continue
    return None, "CSV 인코딩을 해석할 수 없습니다 (UTF-8 / CP949)."


def _read_xlsx_rows(raw):
    ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], None
    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue
        out.append({headers[i]: ("" if (i >= len(r) or r[i] is None) else str(r[i]).strip())
                    for i in range(len(headers))})
    return out, None


def _resolve_headers(keys):
    lower = {str(k).strip().lower(): k for k in keys}
    return {f: lower[f.lower()] for f in IMPORT_HEADER if f.lower() in lower}


def _build_import_host(rows, hmap, existing):
    """같은 hostname 의 행들 → (errors, host)"""
    errors = []
    get = lambda raw, f: (raw.get(hmap[f], "") or "").strip() if f in hmap else None
    first = rows[0][1]
    host = {f: get(first, f) for f in HOST_FIELDS}
    for idx, raw in rows[1:]:
        for f in ("ip", "department"):
            if get(raw, f) and get(raw, f) != host[f]:
                errors.append(f"행 {idx}: {FIELD_LABELS[f]} 가 같은 서버의 다른 행과 다름")
    errs, host = validate_host(host)
    errors += errs

    old_rules = {r["name"].lower(): r for r in (existing or {}).get("rules", [])}
    grouped = {}
    for idx, raw in rows:
        rname = get(raw, "rule_name") or "기본"
        grouped.setdefault(rname, []).append((idx, raw))

    raw_rules = []
    for rname, rrows in grouped.items():
        old = old_rules.get(rname.lower(), {})

        def attr(f, default):
            if f not in hmap:
                return old.get(f, default)
            vals = [get(raw, f) for _, raw in rrows if get(raw, f)]
            return vals[0] if vals else default

        def seq_of(item):
            s = get(item[1], "contact_seq") or ""
            return (int(s) if s.isdigit() else 9999, item[0])

        contacts, seen = [], set()
        for idx, raw in sorted(rrows, key=seq_of):
            phone = normalize_phone(get(raw, "phone") or "")
            if phone in seen:
                continue
            seen.add(phone)
            contacts.append({"name": get(raw, "name") or "", "phone": phone})
        raw_rules.append({
            "id": old.get("id", ""), "name": rname,
            "metrics": _to_list(attr("metrics", ["*"])) or ["*"],
            "keywords": _to_list(attr("keywords", [])),
            "any_level": attr("any_level", "N"), "DAY": attr("DAY", "N"), "NIGHT": attr("NIGHT", "N"),
            "contacts": contacts,
        })
    errs, rules = validate_rules(raw_rules)
    errors += errs
    host["rules"] = rules
    return errors, host


@app.route("/import", methods=["GET", "POST"])
@admin_required
def import_file():
    if request.method == "GET":
        return render_template("import.html", report=None)
    f = request.files.get("file")
    if not f or not f.filename:
        flash("파일을 선택하세요.", "error"); return render_template("import.html", report=None)
    ext = os.path.splitext(f.filename)[1].lower()
    if ext not in ALLOWED_EXT:
        flash("CSV 또는 XLSX 파일만 가능.", "error"); return render_template("import.html", report=None)
    raw = f.read()
    rows, err = (_read_csv_rows(raw) if ext == ".csv" else _read_xlsx_rows(raw))
    if err:
        flash(err, "error"); return render_template("import.html", report=None)
    if not rows:
        flash("데이터 행이 없습니다.", "error"); return render_template("import.html", report=None)
    hmap = _resolve_headers(rows[0].keys())
    missing = [c for c in IMPORT_REQUIRED if c not in hmap]
    if missing:
        flash("필수 열 누락: " + ", ".join(missing) + " (헤더: " + ", ".join(IMPORT_HEADER) + ")", "error")
        return render_template("import.html", report=None)

    report = {"added": 0, "updated": 0, "rejected": [], "warnings": [], "added_hosts": [], "updated_hosts": []}
    groups = {}
    for idx, raw_rec in enumerate(rows, start=2):
        hn = (raw_rec.get(hmap["hostname"], "") or "").strip()
        if not hn:
            report["rejected"].append({"row": str(idx), "hostname": "", "reasons": ["호스트네임 누락"]})
            continue
        groups.setdefault(hn.lower(), []).append((idx, raw_rec))

    with _hosts_lock:
        hosts = load_hosts()
        seen_ip = {}
        for key, grows in groups.items():
            row_label = ",".join(str(i) for i, _ in grows[:5]) + ("…" if len(grows) > 5 else "")
            existing = find_host(hosts, grows[0][1].get(hmap["hostname"], ""))
            errors, host = _build_import_host(grows, hmap, existing)
            if not errors and host["ip"] in seen_ip:
                errors.append(f"파일 내 IP 중복('{seen_ip[host['ip']]}' 와 같음)")
            if not errors:
                owner = find_by_ip(hosts, host["ip"], exclude_hostname=host["hostname"])
                if owner:
                    errors.append(f"IP '{host['ip']}' 가 기존 '{owner['hostname']}' 와 중복")
            if errors:
                report["rejected"].append({"row": row_label, "hostname": host.get("hostname", ""), "reasons": errors})
                continue
            seen_ip[host["ip"]] = host["hostname"]
            for c in contact_conflicts(hosts, host["rules"], exclude_hostname=host["hostname"]):
                report["warnings"].append({"row": row_label, "hostname": host["hostname"], **c})
            if existing:
                hosts = [host if h is existing else h for h in hosts]
                report["updated"] += 1
                report["updated_hosts"].append(host["hostname"])
            else:
                hosts.append(host)
                report["added"] += 1
                report["added_hosts"].append(host["hostname"])
        save_hosts(hosts)
    audit("import", f"추가 {report['added']} / 갱신 {report['updated']} / 거부 {len(report['rejected'])}",
          {"추가": report["added_hosts"], "갱신": report["updated_hosts"], "거부 건수": len(report["rejected"])})
    return render_template("import.html", report=report)


# ===========================================================================
#  규칙 일괄적용 (여러 서버에 같은 규칙 추가/교체/삭제)
# ===========================================================================
@app.route("/monitor")
@login_required
def monitor():
    field, q, pat = _search_args()
    sort = request.args.get("sort", "hostname")
    direction = request.args.get("dir", "asc")
    hosts = sort_hosts(_rule_filter(_search(load_hosts(), field, q), pat), sort, direction)
    rows, page, pages, total = paginate(hosts, _page_arg())
    allh = load_hosts()
    rule_names = sorted({r["name"] for h in allh for r in h.get("rules", [])})
    return render_template("monitor.html", hosts=rows, field=field, q=q, pat=pat,
                           sort=sort, dir=direction, page=page, pages=pages, total=total,
                           rule_names=rule_names, contacts=distinct_contacts(allh),
                           metrics=load_metric_catalog())


@app.route("/monitor/rules", methods=["POST"])
@admin_required
def monitor_rules():
    p = request.get_json(silent=True) or {}
    action = p.get("action", "upsert")
    scope = p.get("scope", "selected")
    if scope == "filtered":
        field = p.get("field", "all") if p.get("field") in SEARCH_FIELDS else "all"
        targets = [h["hostname"] for h in _rule_filter(_search(load_hosts(), field, str(p.get("q", "")).strip()),
                                                       p.get("pat", "all"))]
    else:
        targets = [str(x) for x in p.get("hostnames", [])]
    if not targets:
        return jsonify({"ok": False, "msg": "대상 서버가 없습니다."}), 400

    if action == "delete":
        rname = str(p.get("rule_name", "")).strip()
        if not rname:
            return jsonify({"ok": False, "msg": "삭제할 규칙명을 입력하세요."}), 400
    elif action == "upsert":
        errors, rules = validate_rules([p.get("rule") or {}])
        if errors:
            return jsonify({"ok": False, "msg": "\n".join(errors)}), 400
        rule = rules[0]
        rname = rule["name"]
    else:
        return jsonify({"ok": False, "msg": "알 수 없는 작업"}), 400

    tset = {t.lower() for t in targets}
    changed, skipped = [], []
    with _hosts_lock:
        hosts = load_hosts()
        for h in hosts:
            if h["hostname"].lower() not in tset:
                continue
            rules = h.get("rules", [])
            idx = next((i for i, r in enumerate(rules) if r["name"].lower() == rname.lower()), None)
            if action == "delete":
                if idx is None:
                    continue
                if len(rules) == 1:
                    skipped.append(h["hostname"])      # 마지막 규칙은 삭제 불가
                    continue
                rules.pop(idx)
            else:
                new = dict(rule, id=rules[idx]["id"] if idx is not None else new_rule_id())
                if idx is None:
                    rules.append(new)
                else:
                    rules[idx] = new
            h["rules"] = rules
            changed.append(h["hostname"])
        save_hosts(hosts)
    scope_label = "검색결과 전체" if scope == "filtered" else "선택"
    detail = {"범위": scope_label, "작업": "규칙 삭제" if action == "delete" else "규칙 추가/교체",
              "규칙": rname if action == "delete" else rule_text(rule), "대상 서버": changed}
    if skipped:
        detail["건너뜀(마지막 규칙)"] = skipped
    audit("rules_bulk", f"{scope_label} {len(changed)}대 · {rname}", detail)
    return jsonify({"ok": True, "changed": len(changed), "skipped": skipped})


# ===========================================================================
#  계정 관리
# ===========================================================================
@app.route("/admin/users")
@admin_required
def users_page():
    return render_template("users.html", users=load_users())


@app.route("/admin/users/add", methods=["POST"])
@admin_required
def users_add():
    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    role = request.form.get("role", "viewer")
    if not username or not password:
        flash("아이디와 비밀번호를 입력하세요.", "error"); return redirect(url_for("users_page"))
    if role not in ("admin", "viewer"):
        role = "viewer"
    with _users_lock:
        users = load_users()
        if any(u["username"] == username for u in users):
            flash(f"'{username}' 계정이 이미 존재합니다.", "error"); return redirect(url_for("users_page"))
        users.append({"username": username, "password_hash": generate_password_hash(password),
                      "role": role, "created_at": _now()})
        save_users(users)
    audit("user_add", username, {"role": role})
    flash(f"계정 '{username}' 생성 완료.", "success")
    return redirect(url_for("users_page"))


@app.route("/admin/users/role", methods=["POST"])
@admin_required
def users_role():
    username = request.form.get("username", "")
    role = request.form.get("role", "viewer")
    if role not in ("admin", "viewer"):
        flash("권한 값 오류.", "error"); return redirect(url_for("users_page"))
    with _users_lock:
        users = load_users()
        u = next((x for x in users if x["username"] == username), None)
        if not u:
            flash("계정을 찾을 수 없습니다.", "error"); return redirect(url_for("users_page"))
        if u["role"] == "admin" and role == "viewer" and len([x for x in users if x["role"] == "admin"]) <= 1:
            flash("마지막 관리자 권한은 변경할 수 없습니다.", "error"); return redirect(url_for("users_page"))
        u["role"] = role
        save_users(users)
    audit("user_role", username, {"role": role})
    flash(f"'{username}' 권한을 {role} 로 변경했습니다.", "success")
    return redirect(url_for("users_page"))


@app.route("/admin/users/reset", methods=["POST"])
@admin_required
def users_reset():
    username = request.form.get("username", "")
    password = request.form.get("password", "")
    if not password:
        flash("새 비밀번호를 입력하세요.", "error"); return redirect(url_for("users_page"))
    with _users_lock:
        users = load_users()
        u = next((x for x in users if x["username"] == username), None)
        if not u:
            flash("계정을 찾을 수 없습니다.", "error"); return redirect(url_for("users_page"))
        u["password_hash"] = generate_password_hash(password)
        u.pop("is_default", None)
        save_users(users)
    audit("user_reset", username)
    flash(f"'{username}' 비밀번호를 변경했습니다.", "success")
    return redirect(url_for("users_page"))


@app.route("/admin/users/delete", methods=["POST"])
@admin_required
def users_delete():
    username = request.form.get("username", "")
    if username == session.get("username"):
        flash("본인 계정은 삭제할 수 없습니다.", "error"); return redirect(url_for("users_page"))
    with _users_lock:
        users = load_users()
        u = next((x for x in users if x["username"] == username), None)
        if not u:
            flash("계정을 찾을 수 없습니다.", "error"); return redirect(url_for("users_page"))
        if u["role"] == "admin" and len([x for x in users if x["role"] == "admin"]) <= 1:
            flash("마지막 관리자 계정은 삭제할 수 없습니다.", "error"); return redirect(url_for("users_page"))
        save_users([x for x in users if x["username"] != username])
    audit("user_delete", username)
    flash(f"계정 '{username}' 삭제 완료.", "success")
    return redirect(url_for("users_page"))


@app.route("/change-password", methods=["GET", "POST"])
@login_required
def change_password():
    if request.method == "POST":
        cur, new, new2 = (request.form.get("current", ""), request.form.get("new", ""), request.form.get("new2", ""))
        with _users_lock:
            users = load_users()
            u = next((x for x in users if x["username"] == session["username"]), None)
            if not u or not check_password_hash(u["password_hash"], cur):
                flash("현재 비밀번호가 올바르지 않습니다.", "error"); return render_template("change_password.html")
            if not new or new != new2:
                flash("새 비밀번호가 일치하지 않습니다.", "error"); return render_template("change_password.html")
            u["password_hash"] = generate_password_hash(new)
            u.pop("is_default", None)
            save_users(users)
        session["is_default_pw"] = False
        audit("change_password")
        flash("비밀번호를 변경했습니다.", "success")
        return redirect(url_for("index"))
    return render_template("change_password.html")


# ===========================================================================
#  변경 이력
# ===========================================================================
@app.route("/admin/audit")
@admin_required
def audit_page():
    page = _page_arg()
    entries = []
    if os.path.exists(AUDIT_FILE):
        with open(AUDIT_FILE, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    entries.append(json.loads(line))
                except json.JSONDecodeError:
                    continue
    entries.reverse()
    rows, page, pages, total = paginate(entries, page)
    return render_template("audit.html", entries=rows, page=page, pages=pages,
                           total=total, labels=ACTION_LABELS)


@app.errorhandler(403)
def forbidden(e):
    return render_template("error.html", code=403, msg="권한이 없습니다. (조회 전용 계정)"), 403


@app.errorhandler(404)
def not_found(e):
    return render_template("error.html", code=404, msg="대상을 찾을 수 없습니다."), 404


# ===========================================================================
migrate_existing()
ensure_default_admin()

if __name__ == "__main__":
    app.run(host=SERVER_HOST, port=SERVER_PORT, debug=True)
